"""Build Word (.docx) and Excel (.xlsx) proofreading reports from findings."""
from __future__ import annotations
import datetime as _dt
from typing import List, Dict

PRIORITY_FILL = {"High": "F8CBAD", "Medium": "FFE699", "Low": "D9E1F2"}
PRIORITY_FONT = {"High": "C00000", "Medium": "BF8F00", "Low": "2E5496"}
_HEAD_FILL = "1F3864"
COLS = ["#", "Location", "Language", "Type", "Priority", "Incorrect",
        "Suggestion", "Explanation", "Reference"]


def _ref_label(language: str) -> str:
    return "शब्दकोश ↗" if language in ("Hindi", "Gujarati") else "Dictionary ↗"


def _meta(filename, category, language, n):
    return {
        "File": filename, "Category": category or "General",
        "Language": language, "Total issues": str(n),
        "Generated": _dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


# --------------------------------------------------------------------------- #
def build_docx(findings, out_path, filename, category, language) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def shade(cell, hexcolor):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcolor)
        tcpr.append(sh)

    def add_hyperlink(paragraph, url, text, color="2E5496"):
        part = paragraph.part
        r_id = part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color"); col.set(qn("w:val"), color); rpr.append(col)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rpr.append(sz)
        run.append(rpr)
        t = OxmlElement("w:t"); t.text = text; run.append(t)
        hyperlink.append(run); paragraph._p.append(hyperlink)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Pt(36))

    style = doc.styles["Normal"]
    style.font.name = "Nirmala UI"
    style.element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")
    style.font.size = Pt(10)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Shuddhi — Proofreading Report"); r.bold = True; r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Hindi · Gujarati · English  |  वर्तनी · व्याकरण · संदर्भ")
    rs.font.size = Pt(10); rs.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for k, v in _meta(filename, category, language, len(findings)).items():
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        rk = p.add_run("%s: " % k); rk.bold = True; rk.font.size = Pt(9)
        p.add_run(v).font.size = Pt(9)

    counts = {"High": 0, "Medium": 0, "Low": 0}
    for f in findings:
        counts[f.get("priority", "Medium")] = counts.get(f.get("priority", "Medium"), 0) + 1
    ps = doc.add_paragraph()
    ps.add_run("\nPriority — ").bold = True
    ps.add_run("High: %d   Medium: %d   Low: %d" % (counts["High"], counts["Medium"], counts["Low"]))

    if not findings:
        doc.add_paragraph("\n✅ No issues found.").runs[0].bold = True
        doc.save(out_path); return out_path

    table = doc.add_table(rows=1, cols=len(COLS)); table.style = "Table Grid"
    for ci, name in enumerate(COLS):
        c = table.rows[0].cells[ci]; shade(c, _HEAD_FILL)
        run = c.paragraphs[0].add_run(name); run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); run.font.size = Pt(9)

    for f in findings:
        cells = table.add_row().cells
        vals = [str(f.get("num", "")), f["loc"], f.get("language", ""), f.get("type", ""),
                f.get("priority", ""), f["incorrect"], f["suggestion"],
                f.get("explanation", ""), None]
        for ci, val in enumerate(vals):
            cell = cells[ci]; para = cell.paragraphs[0]
            if COLS[ci] == "Reference" and f.get("reference"):
                add_hyperlink(para, f["reference"], _ref_label(f.get("language", "")))
                continue
            run = para.add_run(val or ""); run.font.size = Pt(9)
            if COLS[ci] == "Incorrect":
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); run.font.strike = True
            elif COLS[ci] == "Suggestion":
                run.font.color.rgb = RGBColor(0x00, 0x86, 0x3B); run.bold = True
            elif COLS[ci] == "Priority":
                pr = f.get("priority", "Medium"); shade(cell, PRIORITY_FILL.get(pr, "FFFFFF"))
                h = PRIORITY_FONT.get(pr, "000000")
                run.font.color.rgb = RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
                run.bold = True

    tip = doc.add_paragraph()
    tip.add_run("\nRed (struck) = incorrect · Green (bold) = suggested correction · "
                "Reference opens a dictionary lookup. Sorted by priority.").italic = True
    doc.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
def build_xlsx(findings, out_path, filename, category, language) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = "Corrections"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Shuddhi — Proofreading Report"; ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    row = 2
    for k, v in _meta(filename, category, language, len(findings)).items():
        ws.cell(row=row, column=1, value=k).font = Font(bold=True, size=9)
        ws.cell(row=row, column=2, value=v).font = Font(size=9)
        row += 1
    header_row = row + 1

    for ci, name in enumerate(COLS, 1):
        c = ws.cell(row=header_row, column=ci, value=name)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border

    for ri, f in enumerate(findings, header_row + 1):
        vals = [f.get("num", ""), f["loc"], f.get("language", ""), f.get("type", ""),
                f.get("priority", ""), f["incorrect"], f["suggestion"],
                f.get("explanation", ""), _ref_label(f.get("language", ""))]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.alignment = Alignment(vertical="top", wrap_text=(COLS[ci - 1] == "Explanation"))
            c.border = border; c.font = Font(size=9)
            name = COLS[ci - 1]
            if name == "Incorrect":
                c.font = Font(size=9, color="C00000", strike=True)
            elif name == "Suggestion":
                c.font = Font(size=9, color="00863B", bold=True)
            elif name == "Reference" and f.get("reference"):
                c.hyperlink = f["reference"]; c.font = Font(size=9, color="2E5496", underline="single")
            elif name == "Priority":
                pr = f.get("priority", "Medium")
                c.fill = PatternFill("solid", fgColor=PRIORITY_FILL.get(pr, "FFFFFF"))
                c.font = Font(size=9, bold=True, color=PRIORITY_FONT.get(pr, "000000"))

    widths = [5, 15, 9, 12, 9, 18, 18, 38, 13]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = "A%d:%s%d" % (header_row, get_column_letter(len(COLS)),
                                       header_row + len(findings))
    wb.save(out_path)
    return out_path
