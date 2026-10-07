"""Text extraction from PDF / DOCX / XLSX / CSV, with optional OCR for scanned PDFs.

Every extractor returns a list of segments:
    {"loc": "<human-readable location>", "text": "<text>"}
so findings can be reported location-by-location (page, paragraph, row, cell).
"""
from __future__ import annotations
import csv
import io
import os
import re
from typing import List, Dict

Segment = Dict[str, str]

# Tesseract language codes per app language choice
_OCR_LANG = {"hi": "hin", "gu": "guj", "en": "eng", "auto": "hin+guj+eng"}

_TESS_DONE = False


def _configure_tesseract(pytesseract) -> None:
    """Point pytesseract at the tesseract binary. On Linux/Docker (Render) it's on
    PATH; on Windows try common install locations."""
    global _TESS_DONE
    if _TESS_DONE:
        return
    _TESS_DONE = True
    import shutil
    if shutil.which("tesseract"):
        return
    for cand in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                 r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                 os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe")):
        if os.path.exists(cand):
            pytesseract.pytesseract.tesseract_cmd = cand
            return


def extract(path: str, ocr: bool = False, ocr_lang: str = "auto") -> List[Segment]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf(path, ocr=ocr, ocr_lang=ocr_lang)
    if ext in (".docx", ".doc"):
        return _docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    if ext in (".csv", ".tsv"):
        return _csv(path)
    if ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as fh:
            return [{"loc": "Line %d" % (i + 1), "text": ln.rstrip("\n")}
                    for i, ln in enumerate(fh) if ln.strip()]
    raise ValueError("Unsupported file type: %s" % ext)


# ---------- PDF ----------
def _pdf(path: str, ocr: bool, ocr_lang: str) -> List[Segment]:
    import pymupdf  # PyMuPDF
    segs: List[Segment] = []
    doc = pymupdf.open(path)
    for i, page in enumerate(doc):
        text = page.get_text().strip()
        garbled = _is_garbled(text)
        # OCR when: forced, OR no text layer, OR the text layer is corrupted
        # (broken font encoding, e.g. PowerPoint→PDF exports).
        if ocr or not text or garbled:
            ocr_text = _ocr_page(page, ocr_lang).strip()
            if ocr_text:
                segs.append({"loc": "Page %d (OCR)" % (i + 1), "text": ocr_text})
                continue
            if text and not garbled:            # OCR unavailable → keep clean text
                segs.append({"loc": "Page %d" % (i + 1), "text": text})
            # garbled text with no OCR result → skip (it would only add noise)
        else:
            segs.append({"loc": "Page %d" % (i + 1), "text": text})
    doc.close()
    return segs


# Stray Latin combining marks that signal a corrupted Devanagari text layer.
_GARBLE_RE = re.compile(r"[̀-ͯǀ-ɏʰ-˿]")
_DEV_RE = re.compile(r"[ऀ-ॿ]")


def _is_garbled(text: str) -> bool:
    if not text:
        return False
    dev = len(_DEV_RE.findall(text))
    bad = len(_GARBLE_RE.findall(text))
    # Latin combining marks never legitimately co-occur with Devanagari, so even
    # a couple are a strong signal of a broken font encoding.
    return dev > 0 and (bad >= 2 or bad >= dev * 0.03)


def _ocr_page(page, ocr_lang: str) -> str:
    """Render a PDF page to an image and OCR it. Requires tesseract installed."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except Exception:
        return ""
    _configure_tesseract(pytesseract)
    pix = page.get_pixmap(dpi=int(os.environ.get("SHUDDHI_OCR_DPI", "220")))
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    lang = _OCR_LANG.get(ocr_lang, _OCR_LANG["auto"])
    try:
        return pytesseract.image_to_string(img, lang=lang)
    except Exception:
        # language data may be missing; fall back to default
        try:
            return pytesseract.image_to_string(img)
        except Exception:
            return ""


# ---------- DOCX ----------
def _docx(path: str) -> List[Segment]:
    from docx import Document
    doc = Document(path)
    segs: List[Segment] = []
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t:
            segs.append({"loc": "Paragraph %d" % (i + 1), "text": t})
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cuff in enumerate(row.cells):
                t = cuff.text.strip()
                if t:
                    segs.append({"loc": "Table %d R%d C%d" % (ti + 1, ri + 1, ci + 1), "text": t})
    return segs


# ---------- XLSX ----------
def _xlsx(path: str) -> List[Segment]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    segs: List[Segment] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip():
                    segs.append({"loc": "%s!%s" % (ws.title, cell.coordinate), "text": v.strip()})
    wb.close()
    return segs


# ---------- CSV / TSV ----------
def _csv(path: str) -> List[Segment]:
    segs: List[Segment] = []
    delim = "\t" if path.lower().endswith(".tsv") else ","
    with open(path, encoding="utf-8", errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=delim)
        for ri, row in enumerate(reader):
            for ci, val in enumerate(row):
                if val and val.strip():
                    segs.append({"loc": "Row %d Col %d" % (ri + 1, ci + 1), "text": val.strip()})
    return segs
